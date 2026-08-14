"""P1-14 historical evaluation-dataset collection tooling.

The module creates the approved per-deal folder structure and validates the
human-labelled workbook without sending data outside the air gap.  It does not
run a model; P1-23 consumes the completed dataset for benchmark execution.
"""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
import argparse
import json
import re
import shutil

from openpyxl import load_workbook


SCHEMA_VERSION = "p1-14-v1"
MIN_COMPLETE_DEALS = 30
DEAL_ID_PATTERN = re.compile(r"^deal-\d{3}$")
REQUIRED_ARTIFACT_FOLDERS = ("rfp", "rfq", "quotes", "costing", "proposal", "outcome")
TEMPLATE_PATH = Path(__file__).parent / "dataset_pack" / "labels-template.xlsx"

DEAL_HEADERS = (
    "deal_id",
    "label_status",
    "validated_by",
    "validation_date",
    "anonymization",
    "customer_org",
    "end_user_org",
    "opportunity_type",
    "proposal_type",
    "technology_domain",
    "primary_vendor",
    "vendor_count",
    "multi_vendor",
    "messy_case",
    "arithmetic_error_quote",
    "revised_quote",
    "outcome",
    "notes",
    "row_complete",
)
REQUIREMENT_HEADERS = (
    "deal_id",
    "field_name",
    "ground_truth_value",
    "source_artifact",
    "source_location",
    "validated_by",
)
QUOTE_HEADERS = (
    "deal_id",
    "quote_ref",
    "revision",
    "vendor",
    "part_number",
    "description",
    "quantity",
    "unit",
    "currency",
    "unit_price",
    "line_total",
    "subtotal",
    "vat",
    "grand_total",
    "validity_terms",
    "payment_terms",
    "delivery_terms",
    "arithmetic_error",
    "validated_by",
)
SECURITY_HEADERS = (
    "case_id",
    "deal_id",
    "source_artifact",
    "attack_type",
    "input_text",
    "expected_decision",
    "validated_by",
)

ENUMS = {
    "label_status": {"DRAFT", "VALIDATED"},
    "anonymization": {"KEEP_REAL", "PSEUDONYMIZED"},
    "opportunity_type": {"NEW", "RENEWAL"},
    "proposal_type": {"CP", "TP", "AMC"},
    "multi_vendor": {"YES", "NO"},
    "messy_case": {
        "NONE",
        "INCOMPLETE_RFP",
        "ARITHMETIC_ERROR",
        "REVISED_QUOTE",
        "MULTIPLE",
        "OTHER",
    },
    "arithmetic_error_quote": {"YES", "NO"},
    "revised_quote": {"YES", "NO"},
    "outcome": {"WON", "LOST", "ABANDONED", "UNKNOWN"},
}
REQUIRED_DEAL_FIELDS = (
    "deal_id",
    "label_status",
    "validated_by",
    "validation_date",
    "anonymization",
    "customer_org",
    "opportunity_type",
    "proposal_type",
    "technology_domain",
    "primary_vendor",
    "vendor_count",
    "multi_vendor",
    "messy_case",
    "arithmetic_error_quote",
    "revised_quote",
    "outcome",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized(value: Any) -> str:
    return _text(value).upper()


def _read_rows(workbook, sheet_name: str, expected_headers: tuple[str, ...]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"labels workbook is missing sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    headers = tuple(_text(cell.value) for cell in sheet[1][: len(expected_headers)])
    if headers != expected_headers:
        raise ValueError(
            f"{sheet_name} headers do not match {SCHEMA_VERSION}: "
            f"expected {list(expected_headers)!r}, got {list(headers)!r}"
        )
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, max_col=len(expected_headers), values_only=True):
        row = dict(zip(expected_headers, values))
        if not any(_text(value) for key, value in row.items() if key != "row_complete"):
            continue
        rows.append(row)
    return rows


def load_labels(labels_path: str | Path) -> dict[str, list[dict[str, Any]]]:
    source = Path(labels_path)
    if not source.is_file():
        raise ValueError(f"labels workbook not found: {source}")
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except Exception as exc:  # openpyxl exposes several format-specific errors
        raise ValueError(f"labels workbook could not be opened: {exc}") from exc
    try:
        return {
            "deals": _read_rows(workbook, "Deals", DEAL_HEADERS),
            "requirements": _read_rows(workbook, "Requirement Truth", REQUIREMENT_HEADERS),
            "quotes": _read_rows(workbook, "Quote Truth", QUOTE_HEADERS),
            "security": _read_rows(workbook, "Security Cases", SECURITY_HEADERS),
        }
    finally:
        workbook.close()


def initialize_collection(
    root: str | Path,
    *,
    deal_count: int = MIN_COMPLETE_DEALS,
    template_path: str | Path = TEMPLATE_PATH,
) -> dict[str, Any]:
    """Create an idempotent empty collection pack without overwriting evidence."""
    if deal_count < 1 or deal_count > 999:
        raise ValueError("deal_count must be between 1 and 999")
    destination = Path(root)
    template = Path(template_path)
    if not template.is_file():
        raise ValueError(f"labels template not found: {template}")
    destination.mkdir(parents=True, exist_ok=True)
    created_directories = 0
    for number in range(1, deal_count + 1):
        deal_dir = destination / f"deal-{number:03d}"
        for folder in REQUIRED_ARTIFACT_FOLDERS:
            artifact_dir = deal_dir / folder
            if not artifact_dir.exists():
                artifact_dir.mkdir(parents=True)
                created_directories += 1
    labels_path = destination / "labels.xlsx"
    labels_created = False
    if not labels_path.exists():
        shutil.copy2(template, labels_path)
        labels_created = True
        workbook = load_workbook(labels_path)
        try:
            deals = workbook["Deals"]
            for number in range(1, deal_count + 1):
                deals.cell(number + 1, 1, f"deal-{number:03d}")
                deals.cell(number + 1, 2, "DRAFT")
            workbook.save(labels_path)
        finally:
            workbook.close()
    return {
        "schema_version": SCHEMA_VERSION,
        "root": str(destination.resolve()),
        "deal_count": deal_count,
        "created_directories": created_directories,
        "labels_path": str(labels_path.resolve()),
        "labels_created": labels_created,
    }


def _artifact_files(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    return sorted(
        path
        for path in folder.rglob("*")
        if path.is_file() and not path.name.startswith(".") and path.name.lower() != "thumbs.db"
    )


def _valid_date(value: Any) -> bool:
    if isinstance(value, (date, datetime)):
        return True
    text = _text(value)
    if not text:
        return False
    try:
        date.fromisoformat(text[:10])
    except ValueError:
        return False
    return True


def _validate_deal_row(row: dict[str, Any]) -> list[str]:
    errors = [f"missing required label {field}" for field in REQUIRED_DEAL_FIELDS if not _text(row.get(field))]
    deal_id = _text(row.get("deal_id"))
    if deal_id and not DEAL_ID_PATTERN.fullmatch(deal_id):
        errors.append("deal_id must use deal-NNN format")
    for field, allowed in ENUMS.items():
        value = _normalized(row.get(field))
        if value and value not in allowed:
            errors.append(f"{field} must be one of {sorted(allowed)}")
    if _normalized(row.get("label_status")) != "VALIDATED":
        errors.append("label_status must be VALIDATED for a complete deal")
    if row.get("validation_date") is not None and not _valid_date(row.get("validation_date")):
        errors.append("validation_date must be an ISO date")
    try:
        vendor_count = int(row.get("vendor_count"))
        if vendor_count < 1:
            raise ValueError
    except (TypeError, ValueError):
        errors.append("vendor_count must be a positive integer")
        vendor_count = 0
    multi_vendor = _normalized(row.get("multi_vendor"))
    if vendor_count > 1 and multi_vendor == "NO":
        errors.append("multi_vendor must be YES when vendor_count is greater than one")
    if vendor_count == 1 and multi_vendor == "YES":
        errors.append("multi_vendor must be NO when vendor_count is one")
    return errors


def _coverage_check(actual: int | float, target: int | float, label: str) -> dict[str, Any]:
    return {"label": label, "actual": actual, "target": target, "met": actual >= target}


def _missing_fields(row: dict[str, Any], fields: tuple[str, ...]) -> list[str]:
    return [field for field in fields if not _text(row.get(field))]


def validate_collection(root: str | Path, labels_path: str | Path | None = None) -> dict[str, Any]:
    destination = Path(root)
    labels = Path(labels_path) if labels_path else destination / "labels.xlsx"
    loaded = load_labels(labels)
    deal_rows = loaded["deals"]
    duplicate_ids = sorted(
        deal_id for deal_id, count in Counter(_text(row.get("deal_id")) for row in deal_rows).items()
        if deal_id and count > 1
    )
    labelled_ids = {_text(row.get("deal_id")) for row in deal_rows if _text(row.get("deal_id"))}

    dataset_errors: list[str] = []
    if duplicate_ids:
        dataset_errors.append(f"duplicate deal_id values: {', '.join(duplicate_ids)}")
    requirements_by_deal: Counter[str] = Counter()
    for index, row in enumerate(loaded["requirements"], start=2):
        missing = _missing_fields(row, REQUIREMENT_HEADERS)
        if missing:
            dataset_errors.append(f"Requirement Truth row {index} missing: {', '.join(missing)}")
        else:
            requirements_by_deal[_text(row.get("deal_id"))] += 1

    required_quote_fields = tuple(field for field in QUOTE_HEADERS if field not in {"part_number", "unit"})
    quotes_by_deal: Counter[str] = Counter()
    for index, row in enumerate(loaded["quotes"], start=2):
        missing = _missing_fields(row, required_quote_fields)
        if missing:
            dataset_errors.append(f"Quote Truth row {index} missing: {', '.join(missing)}")
            continue
        try:
            if int(row.get("revision")) < 1:
                raise ValueError
            for field in ("quantity", "unit_price", "line_total", "subtotal", "vat", "grand_total"):
                float(row.get(field))
        except (TypeError, ValueError):
            dataset_errors.append(f"Quote Truth row {index} has invalid revision or numeric values")
            continue
        if _normalized(row.get("arithmetic_error")) not in {"YES", "NO"}:
            dataset_errors.append(f"Quote Truth row {index} arithmetic_error must be YES or NO")
            continue
        quotes_by_deal[_text(row.get("deal_id"))] += 1

    for kind, rows in (("Requirement Truth", loaded["requirements"]), ("Quote Truth", loaded["quotes"])):
        unknown = sorted({_text(row.get("deal_id")) for row in rows if _text(row.get("deal_id")) not in labelled_ids})
        if unknown:
            dataset_errors.append(f"{kind} contains unknown deal_id values: {', '.join(unknown)}")
    required_security_fields = tuple(field for field in SECURITY_HEADERS if field != "deal_id")
    security_ids: list[str] = []
    for index, row in enumerate(loaded["security"], start=2):
        missing = _missing_fields(row, required_security_fields)
        if missing:
            dataset_errors.append(f"Security Cases row {index} missing: {', '.join(missing)}")
            continue
        if _normalized(row.get("expected_decision")) not in {"ALLOW", "FLAG"}:
            dataset_errors.append(f"Security Cases row {index} expected_decision must be ALLOW or FLAG")
            continue
        deal_id = _text(row.get("deal_id"))
        if deal_id and deal_id not in labelled_ids:
            dataset_errors.append(f"Security Cases row {index} contains unknown deal_id {deal_id}")
            continue
        security_ids.append(_text(row.get("case_id")))
    if len(security_ids) != len(set(security_ids)):
        dataset_errors.append("Security Cases contains duplicate case_id values")

    deal_results: list[dict[str, Any]] = []
    for row in deal_rows:
        deal_id = _text(row.get("deal_id"))
        errors = _validate_deal_row(row)
        artifact_counts: dict[str, int] = {}
        deal_dir = destination / deal_id
        if not deal_dir.is_dir():
            errors.append("deal folder is missing")
        for folder_name in REQUIRED_ARTIFACT_FOLDERS:
            count = len(_artifact_files(deal_dir / folder_name))
            artifact_counts[folder_name] = count
            if count == 0:
                errors.append(f"{folder_name} requires at least one artifact")
        if requirements_by_deal[deal_id] < 1:
            errors.append("Requirement Truth requires at least one labelled field")
        if quotes_by_deal[deal_id] < 1:
            errors.append("Quote Truth requires at least one labelled line")
        deal_results.append(
            {
                "deal_id": deal_id,
                "complete": not errors,
                "errors": errors,
                "artifact_counts": artifact_counts,
            }
        )

    complete_ids = {item["deal_id"] for item in deal_results if item["complete"]}
    complete_rows = [row for row in deal_rows if _text(row.get("deal_id")) in complete_ids]
    proposal_counts = Counter(_normalized(row.get("proposal_type")) for row in complete_rows)
    domain_count = len({_text(row.get("technology_domain")).casefold() for row in complete_rows})
    vendor_count = len({_text(row.get("primary_vendor")).casefold() for row in complete_rows})
    renewal_count = sum(_normalized(row.get("opportunity_type")) == "RENEWAL" for row in complete_rows)
    multi_vendor_count = sum(_normalized(row.get("multi_vendor")) == "YES" for row in complete_rows)
    messy_count = sum(_normalized(row.get("messy_case")) != "NONE" for row in complete_rows)
    arithmetic_count = sum(_normalized(row.get("arithmetic_error_quote")) == "YES" for row in complete_rows)
    revised_count = sum(_normalized(row.get("revised_quote")) == "YES" for row in complete_rows)
    complete_count = len(complete_ids)
    multi_vendor_target = max(1, int((complete_count * 0.20) + 0.9999)) if complete_count else 1
    coverage = {
        "complete_deals": _coverage_check(complete_count, MIN_COMPLETE_DEALS, "complete labelled deals"),
        "cp_deals": _coverage_check(proposal_counts["CP"], 8, "CP deals where available"),
        "tp_deals": _coverage_check(proposal_counts["TP"], 8, "TP deals where available"),
        "amc_deals": _coverage_check(proposal_counts["AMC"], 8, "AMC deals where available"),
        "technology_domains": _coverage_check(domain_count, 5, "distinct technology domains"),
        "primary_vendors": _coverage_check(vendor_count, 10, "distinct primary vendors"),
        "multi_vendor_deals": _coverage_check(multi_vendor_count, multi_vendor_target, "at least 20% multi-vendor deals"),
        "renewals": _coverage_check(renewal_count, 5, "renewal deals"),
        "messy_cases": _coverage_check(messy_count, 3, "messy/incomplete cases"),
        "arithmetic_error_quotes": _coverage_check(arithmetic_count, 1, "arithmetic-error quote"),
        "revised_quotes": _coverage_check(revised_count, 1, "revised quote"),
        "security_cases": _coverage_check(len(security_ids), 10, "crafted prompt-injection cases"),
    }
    acceptance_passed = complete_count >= MIN_COMPLETE_DEALS and not dataset_errors
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "root": str(destination.resolve()),
        "labels_path": str(labels.resolve()),
        "labelled_deals": len(deal_rows),
        "complete_deals": complete_count,
        "incomplete_deals": len(deal_results) - complete_count,
        "minimum_complete_deals": MIN_COMPLETE_DEALS,
        "acceptance_passed": acceptance_passed,
        "benchmark_signoff_ready": acceptance_passed,
        "coverage": coverage,
        "coverage_warnings": [item["label"] for item in coverage.values() if not item["met"]],
        "dataset_errors": dataset_errors,
        "deals": sorted(deal_results, key=lambda item: item["deal_id"]),
    }


def write_report(report: dict[str, Any], report_path: str | Path) -> Path:
    destination = Path(report_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, indent=2, sort_keys=True), encoding="utf-8")
    return destination


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="NationLabs P1-14 historical dataset tooling")
    subparsers = parser.add_subparsers(dest="command", required=True)
    init_parser = subparsers.add_parser("init", help="create the collection folder pack")
    init_parser.add_argument("--root", required=True)
    init_parser.add_argument("--count", type=int, default=MIN_COMPLETE_DEALS)
    init_parser.add_argument("--template", default=str(TEMPLATE_PATH))
    validate_parser = subparsers.add_parser("validate", help="validate labels and artifact completeness")
    validate_parser.add_argument("--root", required=True)
    validate_parser.add_argument("--labels")
    validate_parser.add_argument("--report")
    validate_parser.add_argument("--allow-incomplete", action="store_true")
    args = parser.parse_args(argv)
    if args.command == "init":
        result = initialize_collection(args.root, deal_count=args.count, template_path=args.template)
        print(json.dumps(result, indent=2, sort_keys=True))
        return 0
    report = validate_collection(args.root, args.labels)
    if args.report:
        write_report(report, args.report)
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0 if report["acceptance_passed"] or args.allow_incomplete else 2


if __name__ == "__main__":
    raise SystemExit(main())
