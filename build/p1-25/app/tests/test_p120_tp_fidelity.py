from io import BytesIO
import unittest
import zipfile

from docx import Document
from docx.shared import Inches
from PIL import Image

from integrations.proposal_builder.existing import ExistingProposalBuilderClient
from tp_fidelity import validate_tp_fidelity


def _png_bytes():
    stream = BytesIO()
    image = Image.new("RGB", (64, 64), "white")
    for x in range(32):
        for y in range(64):
            image.putpixel((x, y), (20, 60, 120))
    image.save(stream, format="PNG")
    return stream.getvalue()


def _pdf_with_graphic():
    stream = BytesIO()
    image = Image.new("RGB", (600, 400), "white")
    for x in range(300):
        for y in range(400):
            image.putpixel((x, y), (20, 60, 120))
    image.save(stream, format="PDF")
    return stream.getvalue()


def _docx(
    *,
    shapes=5,
    section_order=("Proposed BOQ", "Commercials", "Acceptance"),
    leak=False,
    client_leak=False,
    money=("75.00", "150.00", "7.50", "157.50"),
):
    document = Document()
    document.add_paragraph("Synthetic approved technical narrative for CUS010.")
    if client_leak:
        document.add_paragraph("Prepared for Northwind Authority.")
    picture = _png_bytes()
    for _ in range(shapes):
        document.add_picture(BytesIO(picture), width=Inches(0.05))
    for heading in section_order:
        document.add_heading(heading, level=1)
        if heading == "Proposed BOQ":
            table = document.add_table(rows=2, cols=4)
            for cell, value in zip(
                table.rows[0].cells, ("S.No", "Part Number", "Description", "Qty")
            ):
                cell.text = value
            for cell, value in zip(
                table.rows[1].cells, ("1", "SELL-1", "Customer appliance", "2")
            ):
                cell.text = value
        elif heading == "Commercials":
            table = document.add_table(rows=4, cols=6)
            for cell, value in zip(
                table.rows[0].cells,
                ("S.No", "Part Number", "Description", "Qty", "Unit Price (AED)", "Total (AED)"),
            ):
                cell.text = value
            for cell, value in zip(
                table.rows[1].cells, ("1", "SELL-1", "Customer appliance", "2", money[0], money[1])
            ):
                cell.text = value
            table.rows[2].cells[-1].text = money[2]
            table.rows[3].cells[-1].text = money[3]
            document.add_paragraph("Payment: 30 days")
            document.add_paragraph("Currency: AED")
            document.add_paragraph("Validity: 30 days")
            document.add_paragraph("Delivery: 2 weeks")
            if leak:
                document.add_paragraph("Vendor cost: 100.00")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _without_first_media(content):
    source = BytesIO(content)
    target = BytesIO()
    with zipfile.ZipFile(source) as archive:
        media = next(name for name in archive.namelist() if name.startswith("word/media/"))
        with zipfile.ZipFile(target, "w") as output:
            for item in archive.infolist():
                if item.filename != media:
                    output.writestr(item, archive.read(item.filename))
    return target.getvalue()


def _frozen_payload():
    return {
        "proposal_type": "TP",
        "template_version": "V1.0",
        "context": {
            "proposal_builder": {
                "proposal_number": "NL-PP-AN-999-26",
                "proposal_version": "V1.0",
                "proposal_date": "2026-08-18",
                "client_name": "CUS010",
                "client_real_name": "Northwind Authority",
                "client_aliases": ["Northwind Authority", "Northwind"],
                "mask_client": True,
                "solution": "Solution",
                "vendor_tp_artifact": {},
                "customer_commercials": {
                    "currency": "AED",
                    "subtotal": "150.00",
                    "vat_rate": "5.00",
                    "vat_amount": "7.50",
                    "grand_total": "157.50",
                    "terms": {
                        "payment": "30 days",
                        "currency": "AED",
                        "validity": "30 days",
                        "delivery": "2 weeks",
                    },
                    "line_items": [
                        {
                            "line_no": 1,
                            "part_number": "SELL-1",
                            "description": "Customer appliance",
                            "quantity": "2",
                            "unit_price": "75.00",
                            "line_total": "150.00",
                        }
                    ],
                    "authority": {
                        "kind": "APPROVED_COSTING_SHEET",
                        "source_sha256": "c" * 64,
                        "approved_by": "finance.user",
                        "approved_at": "2026-08-18T10:00:00+04:00",
                    },
                },
                "rag_provenance": {
                    "status": "APPROVED",
                    "draft_id": 11,
                    "retrieval_id": 12,
                    "reviewed_by": "technical.user",
                    "reviewed_at": "2026-08-18T10:00:00+04:00",
                },
            }
        },
        "commercial": {"native_currency": "USD", "line_items": []},
    }


def _request_payload():
    return {
        "boq": [
            {
                "sno": 1,
                "pn": "SELL-1",
                "desc": "Customer appliance",
                "qty": "2",
                "up": "75.00",
                "total": "150.00",
            }
        ],
        "expectedValue": 157.5,
        "terms": {
            "payment": "30 days",
            "currency": "AED",
            "validity": "30 days",
            "delivery": "2 weeks",
        },
    }


class TpFidelityTests(unittest.TestCase):
    def test_golden_profile_and_new_commercial_order_pass(self):
        source = _docx()
        result = validate_tp_fidelity(
            _docx(),
            request_payload=_request_payload(),
            frozen_payload=_frozen_payload(),
            source_content=source,
            source_filename="vendor.docx",
        )
        self.assertTrue(result["passed"], result)
        self.assertTrue(result["visual_review_required"])
        self.assertEqual(result["profile"], "p1-20.tp-source-aware.v2")
        self.assertEqual(result["metrics"]["drawing_elements"], 5)

    def test_wrong_order_graphic_loss_identity_leak_and_internal_cost_fail_closed(self):
        result = validate_tp_fidelity(
            _docx(
                shapes=3,
                section_order=("Proposed BOQ", "Acceptance", "Commercials"),
                leak=True,
                client_leak=True,
            ),
            request_payload=_request_payload(),
            frozen_payload=_frozen_payload(),
            source_content=_docx(shapes=5),
            source_filename="vendor.docx",
        )
        self.assertFalse(result["passed"])
        failed = {item["name"] for item in result["checks"] if not item["passed"]}
        self.assertIn("tp_section_order", failed)
        self.assertIn("source_graphics_preserved", failed)
        self.assertIn("masked_client_identity", failed)
        self.assertIn("no_internal_commercial_labels", failed)

    def test_formatted_money_matches_frozen_numeric_facts(self):
        request = _request_payload()
        request["boq"][0].update({"qty": "2", "up": "120000.00", "total": "240000.00"})
        request["expectedValue"] = 252000.00
        result = validate_tp_fidelity(
            _docx(money=("120,000.00", "240,000.00", "12,000.00", "252,000.00")),
            request_payload=request,
            frozen_payload=_frozen_payload(),
            source_content=_docx(),
            source_filename="vendor.docx",
        )
        check = next(item for item in result["checks"] if item["name"] == "approved_commercial_facts_present")
        self.assertTrue(check["passed"], check)

    def test_missing_image_relationship_target_fails_closed(self):
        result = validate_tp_fidelity(
            _without_first_media(_docx()),
            request_payload=_request_payload(),
            frozen_payload=_frozen_payload(),
            source_content=_docx(),
            source_filename="vendor.docx",
        )
        self.assertFalse(result["passed"])
        failed = {item["name"] for item in result["checks"] if not item["passed"]}
        # python-docx itself rejects this package before the relationship-level
        # report runs, which is the earliest and strongest fail-closed outcome.
        self.assertIn("valid_docx_package", failed)

    def test_missing_source_evidence_fails_closed(self):
        result = validate_tp_fidelity(
            _docx(), request_payload=_request_payload(), frozen_payload=_frozen_payload()
        )
        self.assertFalse(result["passed"])
        failed = {item["name"] for item in result["checks"] if not item["passed"]}
        self.assertIn("source_graphics_preserved", failed)

    def test_pdf_source_graphics_are_measured_without_a_golden_shape_floor(self):
        result = validate_tp_fidelity(
            _docx(shapes=1),
            request_payload=_request_payload(),
            frozen_payload=_frozen_payload(),
            source_content=_pdf_with_graphic(),
            source_filename="vendor.pdf",
        )
        check = next(item for item in result["checks"] if item["name"] == "source_graphics_preserved")
        self.assertTrue(check["passed"], check)
        self.assertEqual(result["metrics"]["source_graphics"], 1)

    def test_tp_maps_only_approved_customer_selling_prices(self):
        payload = _frozen_payload()
        payload["context"]["proposal_builder"]["vendor_tp_artifact"] = {
            "artifact_ref": "file:///tmp/vendor.docx",
            "sha256": "d" * 64,
        }
        import hashlib
        import json

        payload_hash = hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
        ).hexdigest()
        mapped = ExistingProposalBuilderClient._builder_payload(
            proposal_type="TP",
            template_version="V1.0",
            payload=payload,
            payload_hash=payload_hash,
        )
        self.assertEqual(mapped["boq"][0]["up"], "75.00")
        self.assertEqual(mapped["boq"][0]["total"], "150.00")
        self.assertEqual(mapped["expectedValue"], 157.5)
        self.assertEqual(mapped["client"], "CUS010")
        self.assertEqual(mapped["clientRealName"], "Northwind Authority")
        self.assertEqual(mapped["clientAliases"], "Northwind Authority\nNorthwind")
        self.assertTrue(mapped["maskClient"])
        self.assertNotIn("cost", str(mapped).casefold())

    def test_tp_without_explicit_client_masking_is_rejected(self):
        import hashlib
        import json

        payload = _frozen_payload()
        payload["context"]["proposal_builder"]["mask_client"] = False
        payload_hash = hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
        ).hexdigest()
        with self.assertRaisesRegex(ValueError, "mask_client=true"):
            ExistingProposalBuilderClient._builder_payload(
                proposal_type="TP",
                template_version="V1.0",
                payload=payload,
                payload_hash=payload_hash,
            )

    def test_tp_without_approved_costing_sheet_is_rejected(self):
        payload = _frozen_payload()
        del payload["context"]["proposal_builder"]["customer_commercials"]
        import hashlib
        import json

        payload_hash = hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
        ).hexdigest()
        with self.assertRaisesRegex(ValueError, "approved costing sheet"):
            ExistingProposalBuilderClient._builder_payload(
                proposal_type="TP",
                template_version="V1.0",
                payload=payload,
                payload_hash=payload_hash,
            )

    def test_tp_without_approved_rag_provenance_is_rejected_before_build(self):
        import hashlib
        import json

        payload = _frozen_payload()
        del payload["context"]["proposal_builder"]["rag_provenance"]
        payload_hash = hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
        ).hexdigest()
        with self.assertRaisesRegex(ValueError, "approved RAG provenance"):
            ExistingProposalBuilderClient._builder_payload(
                proposal_type="TP",
                template_version="V1.0",
                payload=payload,
                payload_hash=payload_hash,
            )

    def test_tp_rejects_bad_vat_and_internal_cost_fields(self):
        import hashlib
        import json

        for mutation, error in (("vat", "VAT amount"), ("cost", "internal fields")):
            with self.subTest(mutation=mutation):
                payload = _frozen_payload()
                snapshot = payload["context"]["proposal_builder"]["customer_commercials"]
                if mutation == "vat":
                    snapshot["vat_amount"] = "8.00"
                    snapshot["grand_total"] = "158.00"
                else:
                    snapshot["line_items"][0]["vendor_cost"] = "40.00"
                payload_hash = hashlib.sha256(
                    json.dumps(
                        payload,
                        ensure_ascii=False,
                        separators=(",", ":"),
                        sort_keys=True,
                    ).encode()
                ).hexdigest()
                with self.assertRaisesRegex(ValueError, error):
                    ExistingProposalBuilderClient._builder_payload(
                        proposal_type="TP",
                        template_version="V1.0",
                        payload=payload,
                        payload_hash=payload_hash,
                    )


if __name__ == "__main__":
    unittest.main()
