"""Deterministic P1-20 structural fidelity checks for generated TP DOCX files.

This module validates Proposal Builder output.  It does not generate or edit
Word documents and therefore keeps the frozen Builder as the sole document
owner.  Visual raster comparison remains a document-worker responsibility.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
import posixpath
import re
from typing import Any
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
import pdfplumber
from PIL import Image, ImageStat, UnidentifiedImageError


PROFILE_VERSION = "p1-20.tp-source-aware.v2"
_REQUIRED_HEADINGS = ("proposed boq", "commercials", "acceptance")
_FORBIDDEN_CUSTOMER_LABELS = (
    "vendor cost",
    "buy price",
    "purchase cost",
    "internal cost",
    "gross margin",
    "margin %",
    "markup %",
)


def _normalized(value: Any) -> str:
    return " ".join(str(value or "").replace("\ufffd", " ").split()).casefold()


def _heading_name(value: str) -> str:
    value = re.sub(r"^\s*\d+(?:\.\d+)*\s*[.)-]?\s*", "", value)
    return _normalized(value)


def _check(name: str, passed: bool, details: Any) -> dict[str, Any]:
    return {"name": name, "passed": bool(passed), "details": details}


def _decimal(value: Any) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, TypeError, ValueError):
        return None


def _numeric_tokens(value: str) -> set[Decimal]:
    tokens: set[Decimal] = set()
    pattern = r"(?<![\w.])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\w.])"
    for match in re.finditer(pattern, value):
        number = _decimal(match.group(0))
        if number is not None:
            tokens.add(number)
    return tokens


def _all_docx_text(archive: zipfile.ZipFile) -> str:
    """Extract text from every Word story, including headers and text boxes."""
    values: list[str] = []
    for name in archive.namelist():
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        try:
            root = ET.fromstring(archive.read(name))
        except ET.ParseError:
            continue
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1] in {"t", "instrText"} and node.text:
                values.append(node.text)
    return "\n".join(values)


def _relationship_source(rel_name: str) -> str | None:
    marker = "/_rels/"
    if marker not in rel_name or not rel_name.endswith(".rels"):
        return None
    parent, filename = rel_name.split(marker, 1)
    return posixpath.join(parent, filename[:-5])


def _docx_image_report(archive: zipfile.ZipFile) -> dict[str, Any]:
    """Verify that every internal image relationship resolves to usable media."""
    names = set(archive.namelist())
    targets: list[str] = []
    missing: list[str] = []
    corrupt: list[str] = []
    blank: list[str] = []
    image_relationships = 0

    for rel_name in sorted(name for name in names if name.endswith(".rels")):
        source = _relationship_source(rel_name)
        if source is None:
            continue
        try:
            root = ET.fromstring(archive.read(rel_name))
        except ET.ParseError:
            corrupt.append(rel_name)
            continue
        for relationship in root:
            rel_type = relationship.attrib.get("Type", "")
            if not rel_type.endswith("/image") or relationship.attrib.get("TargetMode") == "External":
                continue
            image_relationships += 1
            target = relationship.attrib.get("Target", "")
            resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source), target))
            targets.append(resolved)
            if resolved not in names:
                missing.append(resolved)

    dimensions: dict[str, list[int]] = {}
    for target in sorted(set(targets) - set(missing)):
        content = archive.read(target)
        if not content:
            corrupt.append(target)
            continue
        try:
            with Image.open(BytesIO(content)) as image:
                image.load()
                dimensions[target] = [int(image.width), int(image.height)]
                extrema = ImageStat.Stat(image.convert("L")).extrema[0]
                if image.width * image.height >= 4096 and extrema[0] == extrema[1]:
                    blank.append(target)
        except (UnidentifiedImageError, OSError, ValueError):
            corrupt.append(target)

    drawing_elements = 0
    main_document_drawings = 0
    for name in names:
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        try:
            root = ET.fromstring(archive.read(name))
        except ET.ParseError:
            continue
        part_drawings = sum(
            1 for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "drawing"
        )
        drawing_elements += part_drawings
        if name == "word/document.xml":
            main_document_drawings = part_drawings

    return {
        "image_relationships": image_relationships,
        "distinct_media": len(set(targets)),
        "drawing_elements": drawing_elements,
        "main_document_drawings": main_document_drawings,
        "missing_targets": sorted(set(missing)),
        "corrupt_media": sorted(set(corrupt)),
        "blank_media": sorted(set(blank)),
        "dimensions": dimensions,
    }


def _source_graphics(content: bytes | None, filename: str | None) -> dict[str, Any]:
    if not content:
        return {"kind": "unavailable", "graphics": None}
    suffix = (filename or "").lower().rsplit(".", 1)[-1]
    if suffix == "docx":
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                report = _docx_image_report(archive)
            return {"kind": "docx", "graphics": report["main_document_drawings"]}
        except (OSError, zipfile.BadZipFile):
            return {"kind": "docx", "graphics": None, "error": "source DOCX could not be inspected"}
    if suffix == "pdf":
        try:
            graphics = 0
            with pdfplumber.open(BytesIO(content)) as pdf:
                for page in pdf.pages:
                    page_area = max(float(page.width) * float(page.height), 1.0)
                    graphics += sum(
                        1
                        for item in page.images
                        if abs(float(item.get("width") or 0) * float(item.get("height") or 0))
                        >= max(4096.0, page_area * 0.005)
                    )
            return {"kind": "pdf", "graphics": graphics}
        except Exception as exc:
            return {"kind": "pdf", "graphics": None, "error": str(exc)[:300]}
    return {"kind": suffix or "unknown", "graphics": None}


def _identity_contract(frozen_payload: dict[str, Any]) -> dict[str, Any]:
    context = frozen_payload.get("context")
    builder = context.get("proposal_builder", context) if isinstance(context, dict) else {}
    aliases = builder.get("client_aliases") if isinstance(builder, dict) else []
    if isinstance(aliases, str):
        aliases = [item.strip() for item in aliases.splitlines() if item.strip()]
    elif not isinstance(aliases, (list, tuple)):
        aliases = []
    real_name = str(builder.get("client_real_name") or "").strip() if isinstance(builder, dict) else ""
    values = ([real_name] if real_name else []) + [str(item).strip() for item in aliases if str(item).strip()]
    code = str(builder.get("client_name") or builder.get("client") or "").strip() if isinstance(builder, dict) else ""
    unique = []
    seen = set()
    for item in values:
        key = _normalized(item)
        if not key or key == _normalized(code) or key in seen:
            continue
        seen.add(key)
        unique.append(item)
    return {
        "mask_client": bool(builder.get("mask_client")) if isinstance(builder, dict) else False,
        "client_code": code,
        "aliases": unique,
    }


def validate_tp_fidelity(
    content: bytes,
    *,
    request_payload: dict[str, Any],
    frozen_payload: dict[str, Any],
    source_content: bytes | None = None,
    source_filename: str | None = None,
) -> dict[str, Any]:
    """Return a fail-closed, JSON-serializable validation report."""
    checks: list[dict[str, Any]] = []
    try:
        document = Document(BytesIO(content))
        with zipfile.ZipFile(BytesIO(content)) as archive:
            all_text = _all_docx_text(archive)
            image_report = _docx_image_report(archive)
    except Exception as exc:
        return {
            "passed": False,
            "profile": PROFILE_VERSION,
            "checks": [_check("valid_docx_package", False, str(exc))],
            "metrics": {},
        }

    checks.append(_check("valid_docx_package", True, "DOCX opened successfully"))
    headings = [
        _heading_name(paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.text.strip() and paragraph.style.name.casefold().startswith("heading")
    ]
    positions: list[int] = []
    for expected in _REQUIRED_HEADINGS:
        matches = [index for index, heading in enumerate(headings) if heading == expected]
        positions.append(matches[0] if len(matches) == 1 else -1)
    ordered = all(position >= 0 for position in positions) and positions == sorted(positions)
    adjacent = ordered and positions[1] == positions[0] + 1 and positions[2] == positions[1] + 1
    checks.append(
        _check(
            "tp_section_order",
            ordered and adjacent,
            {
                "required": list(_REQUIRED_HEADINGS),
                "positions": positions,
                "headings": headings,
                "commercials_immediately_after_boq": adjacent,
            },
        )
    )

    table_headers = []
    for table in document.tables:
        if table.rows:
            table_headers.append([_normalized(cell.text) for cell in table.rows[0].cells])
    boq_table = any(
        all(label in headers for label in ("s.no", "part number", "description", "qty"))
        for headers in table_headers
    )
    commercial_table = any(
        all(
            any(required in header for header in headers)
            for required in ("part number", "description", "qty", "unit price", "total")
        )
        for headers in table_headers
    )
    checks.append(
        _check(
            "boq_and_commercial_tables",
            boq_table and commercial_table,
            {"boq_table": boq_table, "commercial_table": commercial_table},
        )
    )

    normalized_text = _normalized(all_text)
    numeric_tokens = _numeric_tokens(all_text)
    missing_facts: list[str] = []
    for item in request_payload.get("boq") or []:
        for field in ("pn", "desc", "qty", "up", "total"):
            value = _normalized(item.get(field))
            if field in {"qty", "up", "total"}:
                present = _decimal(item.get(field)) in numeric_tokens
            else:
                present = not value or value in normalized_text
            if value and not present:
                missing_facts.append(f"boq[{item.get('sno')}].{field}={item.get(field)}")
    expected_value = _decimal(request_payload.get("expectedValue"))
    if expected_value is not None and expected_value not in numeric_tokens:
        missing_facts.append(f"expectedValue={request_payload.get('expectedValue')}")
    terms = request_payload.get("terms") if isinstance(request_payload.get("terms"), dict) else {}
    for field in ("payment", "currency", "validity", "delivery"):
        value = _normalized(terms.get(field))
        if value and value not in normalized_text:
            missing_facts.append(f"terms.{field}={terms.get(field)}")
    checks.append(
        _check(
            "approved_commercial_facts_present",
            not missing_facts,
            {"missing": missing_facts},
        )
    )

    identity = _identity_contract(frozen_payload)
    leaked_aliases = [
        f"alias #{index}"
        for index, alias in enumerate(identity["aliases"], 1)
        if _normalized(alias) in normalized_text
    ]
    code_present = bool(identity["client_code"] and _normalized(identity["client_code"]) in normalized_text)
    identity_valid = bool(
        identity["mask_client"]
        and identity["aliases"]
        and code_present
        and not leaked_aliases
    )
    checks.append(
        _check(
            "masked_client_identity",
            identity_valid,
            {
                "mask_client": identity["mask_client"],
                "client_code_present": code_present,
                "alias_count": len(identity["aliases"]),
                "leaked_aliases": leaked_aliases,
            },
        )
    )

    leaked_labels = [label for label in _FORBIDDEN_CUSTOMER_LABELS if label in normalized_text]
    checks.append(
        _check(
            "no_internal_commercial_labels",
            not leaked_labels,
            {"forbidden_labels_found": leaked_labels},
        )
    )

    context = frozen_payload.get("context")
    builder_context = context.get("proposal_builder", context) if isinstance(context, dict) else {}
    rag = builder_context.get("rag_provenance") if isinstance(builder_context, dict) else None
    rag_valid = bool(
        isinstance(rag, dict)
        and _normalized(rag.get("status")) == "approved"
        and rag.get("draft_id")
        and rag.get("retrieval_id")
        and str(rag.get("reviewed_by") or "").strip()
        and str(rag.get("reviewed_at") or "").strip()
    )
    checks.append(
        _check(
            "approved_rag_provenance",
            rag_valid,
            "approved draft/retrieval/reviewer provenance is required",
        )
    )

    relationship_integrity = not (
        image_report["missing_targets"]
        or image_report["corrupt_media"]
        or image_report["blank_media"]
    )
    checks.append(
        _check(
            "embedded_object_relationship_integrity",
            relationship_integrity,
            image_report,
        )
    )

    source_report = _source_graphics(source_content, source_filename)
    source_count = source_report.get("graphics")
    # Header/footer logos are template branding, not preserved vendor figures;
    # compare the source only with drawings in the main document story.
    output_count = image_report["main_document_drawings"]
    source_preserved = source_count is not None and (
        source_count == 0 or output_count >= source_count
    )
    checks.append(
        _check(
            "source_graphics_preserved",
            source_preserved,
            {
                "source": source_report,
                "output_drawing_elements": output_count,
                "rule": "output drawings must meet or exceed detected source graphics",
            },
        )
    )

    metrics = {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "media_parts": image_report["distinct_media"],
        "image_relationships": image_report["image_relationships"],
        "drawing_elements": image_report["drawing_elements"],
        "main_document_drawings": image_report["main_document_drawings"],
        "source_graphics": source_count,
        "sections": len(document.sections),
    }
    return {
        "passed": all(check["passed"] for check in checks),
        "profile": PROFILE_VERSION,
        "checks": checks,
        "metrics": metrics,
        "visual_review_required": True,
    }
